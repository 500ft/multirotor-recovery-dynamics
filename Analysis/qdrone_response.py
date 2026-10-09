"""Descriptive response summary of the registered QDrone2 development recording."""
import argparse
import hashlib
import json
from pathlib import Path
import subprocess
import time

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
RECORD = ROOT / 'Data/qdrone-response'
HEADERS = ['Time (s)', 'zd (m)', 'zw (m)', 'Battery Voltage (V)',
           'm1 (V)', 'm2 (V)', 'm3 (V)', 'm4 (V)', 'xw (m)', 'yw (m)',
           'roll (rad)', 'pitch (rad)', 'yaw (rad)']


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_recording(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        if book.sheetnames != ['Sheet1']:
            raise ValueError('unexpected source sheets')
        rows = book['Sheet1'].iter_rows(values_only=True)
        if list(next(rows)) != HEADERS:
            raise ValueError('source columns or units differ from inspected header')
        # Read literal source values, never evaluate spreadsheet formulas or
        # silently substitute their potentially stale cached results.
        data = np.asarray(list(rows), dtype=float)
    finally:
        book.close()
    if data.ndim != 2 or data.shape[1] != len(HEADERS):
        raise ValueError('unexpected recording shape')
    return data


def summarize(data, protocol):
    t, reference, altitude, voltage = data[:, :4].T
    if len(t) < 2 or not np.isfinite(t).all() or not (np.diff(t) > 0).all():
        raise ValueError('finite strictly increasing timestamps required')
    if not np.isfinite(reference).all():
        raise ValueError('cannot identify reference changes with missing command')
    starts = np.flatnonzero(abs(np.diff(reference)) > protocol['step_threshold_m']) + 1
    events = []
    for ordinal, start in enumerate(starts):
        end = t[start] + protocol['horizon_s']
        next_start = int(starts[ordinal + 1]) if ordinal + 1 < len(starts) else None
        event = dict(step=ordinal + 1, source_row_zero_based=int(start), time_s=float(t[start]),
            target_m=float(reference[start]), change_m=float(reference[start]-reference[start-1]),
            direction='up' if reference[start] > reference[start-1] else 'down',
            precommand_voltage_v=float(voltage[start-1]) if np.isfinite(voltage[start-1]) else None,
            precommand_voltage_time_s=float(t[start-1]), horizon_s=protocol['horizon_s'])
        events.append(event)
        if end > t[-1] + 1e-12:
            event.update(status='incomplete', reason='recording_ended'); continue
        if next_start is not None and t[next_start] < end - 1e-12:
            event.update(status='incomplete', reason='next_command_inside_horizon'); continue
        # Use observed timestamps, adding just the prescribed endpoint. This
        # interpolation scores an observed response, not a causal predictor.
        stop = int(np.searchsorted(t, end, side='left'))
        stop = min(stop, len(t)-1)
        if not np.isfinite(data[start:stop+1, 2:4]).all():
            event.update(status='incomplete', reason='nonfinite_response'); continue
        ts = np.r_[t[start:stop], end]
        z = np.r_[altitude[start:stop], np.interp(end, t, altitude)]
        v = np.r_[voltage[start:stop], np.interp(end, t, voltage)]
        error = z - reference[start]
        duration = ts[-1]-ts[0]
        mean = lambda x: float(np.trapezoid(x, ts)/duration)
        event.update(status='complete', bias_m=mean(error), mae_m=mean(abs(error)),
            rmse_m=float(np.sqrt(mean(error**2))), mean_voltage_v=mean(v),
            samples_used=int(stop-start+1), maximum_interval_s=float(np.diff(ts).max()))
    dt = np.diff(t)
    return dict(rows=len(t), time_start_s=float(t[0]), time_end_s=float(t[-1]),
        dt_min_s=float(dt.min()), dt_median_s=float(np.median(dt)), dt_max_s=float(dt.max()),
        columns=HEADERS, detected_steps=len(events), complete_steps=sum(e['status']=='complete' for e in events),
        initial_steady_fragment='Not scored as a commanded response.',
        final_recorded_voltage_v=float(voltage[-1]), events=events)


def plot(data, result, path):
    # Keep the existing analysis CLI compatible; the renderer also accepts the
    # committed result directly, so visual changes need not rewrite evidence.
    from Analysis.plot_qdrone_response import plot as render
    render(data, result, path)


def run(cache, output, figure, fetch=False):
    started = time.perf_counter()
    sources = json.loads((RECORD/'sources.json').read_text())
    protocol = json.loads((RECORD/'protocol.json').read_text())
    for item in sources['files']:
        path = cache/item['name']
        if fetch and not path.exists():
            cache.mkdir(parents=True, exist_ok=True)
            subprocess.run(['curl', '-fsSL', item['url'], '-o', str(path)], check=True)
        if digest(path) != item['sha256']:
            raise ValueError(f'input hash mismatch: {item["name"]}')
    data = read_recording(cache/protocol['selected_file'])
    result = summarize(data, protocol)
    result.update(dataset='QDrone2', role=protocol['role'], source_sha256=protocol['selected_sha256'],
        protocol_sha256=digest(RECORD/'protocol.json'), code_sha256=digest(Path(__file__)),
        source_record='Data/qdrone-response/sources.json', protocol_horizon_s=protocol['horizon_s'],
        selected_recordings=1, independently_identified_packs=None, archive_independent_run_count=None,
        independence='Repeated commands share one discharge. Archive figure panels do not establish independent runs or pack identities.',
        endpoint='Continuous altitude tracking under original MPC. Source describes a voltage-based stop rule; recording end alone is not a labeled recovery failure.',
        warning_performance='Not evaluated; no independent recovery outcomes or fitted predictor.',
        runtime_s=time.perf_counter()-started)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, indent=2, allow_nan=False)+'\n')
    if figure:
        plot(data, result, figure)
    print(json.dumps({k:result[k] for k in ['rows', 'detected_steps', 'complete_steps', 'runtime_s']}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--cache', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--figure', type=Path)
    parser.add_argument('--fetch', action='store_true')
    args = parser.parse_args()
    run(args.cache, args.output, args.figure, args.fetch)
