import ast, csv, dataclasses, json, math, pathlib, re, sys, hashlib
import numpy as np
import openpyxl

BASE = pathlib.Path(__file__).parent
ROOT = BASE / 'source'
sys.path.insert(0, str(ROOT))
SHA = 'fbc4c878a99f68650c70670e7ee3721a6a24e67d'
URL = 'https://github.com/500ft/multirotor-recovery-dynamics/blob/' + SHA + '/'
tabs = []
def link(p): return URL + str(p).replace(' ', '%20')
def add(name, rows, source='', note='', header=1, formulas=False):
    width = max(map(len, rows), default=1)
    rows = [r + [None]*(width-len(r)) for r in rows]
    tab = dict(id=100+len(tabs), name=name, rows=rows, width=width, source=source, note=note, header=header, formulas=formulas)
    tabs.append(tab)
    return tab
def flatten(d, prefix=''):
    out={}
    for k,v in d.items():
        key=prefix+k
        if isinstance(v,dict): out.update(flatten(v,key+'.'))
        elif isinstance(v,list):out[key]=json.dumps(v,ensure_ascii=False)
        else:out[key]=v
    return out
def records(name, recs, source, note=''):
    flat=[flatten(r) for r in recs]
    keys=list(dict.fromkeys(k for r in flat for k in r))
    return add(name,[keys]+[[r.get(k) for k in keys] for r in flat],source,note)
def readj(p):return json.loads((ROOT/p).read_text())
def num(s):
    if re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?(?:[eE][+-]?\d+)?',s or ''):
        return float(s) if '.' in s or 'e' in s.lower() else int(s)
    return s
guide=add('Start Here',[['Topic','Details'],['Project','Multirotor Recovery Dynamics — calculations and simulation archive'],['Source snapshot',SHA],['Imported on','2026-09-28'],['Scope','All six committed simulation JSON files; engineering CSVs; component studies; equations, model source, plot sources and a reproduced release time series.'],['Evidence','Simulation outputs describe the historical design profile. They are not measurements of the Veeniix V995.'],['V995','Owner-estimated mass around 50 g. Do not substitute the historical 135.66 g design budget for a measured V995 mass.'],['Raw data','Original JSON fields and CSV values are preserved. Nested objects become dotted columns. JSON arrays in metadata remain JSON text. Blank values remain unavailable, not zero.'],['Live calculations','Sizing, guard, mixer, statistics and research workbook formulas recalculate in Sheets. Imported simulation outputs are snapshots; editing inputs does not rerun Python dynamics.'],['Scenarios','Study A and A2 are release_startup. Scenario Trials holds the newer three-arm exploratory diagnostic. Do not pool these as equivalent experiments.'],['Statistics','Bounds describe Monte Carlo sampling uncertainty conditional on the model. Zero successes does not establish equivalence or physical impossibility.'],['Historical policy','Package Rankings preserves the source policy key as a historical package/configuration ranking, not a verified runtime switching policy.'],['Missing data','Full per-step traces for historical Monte Carlo and survivable-set sweeps were not committed. Do not reconstruct observations from summary statistics.'],['Reproduction','Release Trace is regenerated from pinned source with nominal parameters, 2 rad/s, 60 degrees, release_startup. The three stored demo outputs were matched exactly.'],['Research archives','CS tabs preserve the first component workbook and its formulas. Component Selection preserves every nonempty row of the second workbook, including section headers. Catalog data retains its original dates.'],['Refresh','Refresh source tables from a new pinned commit, then rebuild charts and recalculate. This workbook is a snapshot, not a live GitHub sync.']])
charts=add('Charts',[['Chart','Meaning'],['Release plots','Reproduced nominal release trace; historical design and placeholder torque authority.'],['Study plots','Frozen touchdown proxy, not physical flight qualification.'],['Diagnostic plots','Limited exploratory prescribed-upset cohort; three scenario arms, 12 paired draws per cell/package/class.']])
sources=add('Source Index',[['Tab','Source path','Source URL','Rows excluding header','Notes']])

csvs={
'Mass Budget':'Engineering Data/mass_budget.csv','Power Budget':'Engineering Data/power_budget.csv',
'Estimates':'Engineering Data/estimates.csv','Component Verification':'Engineering Data/component_verification.csv',
'Hardware Interfaces':'Engineering Data/hardware_interfaces.csv','Instrumentation':'Engineering Data/instrumentation.csv',
'Requirements':'Engineering Data/requirements.csv','BOM':'Design Report/BOM.csv',
'Purchase Checklist':'Engineering Data/purchase_checklist.csv','FMEA':'Engineering Data/fmea.csv',
'CAD Parameters':'cad/bench/parameters.csv','CAD Input Requests':'cad/bench/input-requests.csv',
'Wire Data':'Engineering Data/onshape_wire_data.csv','Hardware Inventory':'evidence/week-2026-09-19/hardware-inventory.csv',
'Interface Observations':'evidence/week-2026-09-19/interface-observations.csv',
'CAD Mass Template':'Engineering Data/cad_mass_properties_template.csv'}
for name,p in csvs.items():
    rows=list(csv.reader((ROOT/p).open()))
    add(name,[rows[0]]+[[num(v) for v in r] for r in rows[1:]],p,'Source status is retained; CONFIRMED catalog specifications are not bench measurements. Templates contain no measurements.' if 'Template' in name or 'Observations' in name else 'Imported without changing source evidence status.')

meta=[]
def metadata(source, obj):
    for k,v in flatten(obj).items():meta.append({'source':source,'field':k,'value':v,'value_type':'null' if v is None else type(v).__name__})
for variant in ['a','a2']:
    p='Data/survivable_set_results'+('' if variant=='a' else '_a2')+'.json';d=readj(p)
    records('Study '+variant.upper(),[dict(stage=stage,**r) for stage in ['exploratory','primary'] for r in d[stage]],p,'Historical release_startup. Assumed landing thresholds; simulation only.')
    records('Paired '+variant.upper(),d['paired_mechanism_vs_realloc'],p,'Paired mechanism versus reallocation. Null interval bounds remain blank.')
    if variant=='a':
        records('A vs A2 Paired',d['paired_variant_comparison'],p,'Controller and drag change together; not an isolated causal effect.')
        rankings=[];kills=[]
    rankings += [dict(study=variant,**r) for r in d['policy']]
    kills += [dict(study=variant,**r) for r in d['kill_criterion']['comparisons']]
    metadata(p,{k:v for k,v in d.items() if k not in ['exploratory','primary','paired_mechanism_vs_realloc','paired_variant_comparison','policy','kill_criterion']})
    metadata(p,{'kill_criterion':{k:v for k,v in d['kill_criterion'].items() if k!='comparisons'}})
records('Package Rankings',rankings,'Data/survivable_set_results*.json','Source policy records; configuration ranking only.')
records('Primary Comparisons',kills,'Data/survivable_set_results*.json','Historical kill-criterion outputs, not a non-inferiority test.')
p='Data/scenario-diagnostic/scenario_diagnostic.json';d=readj(p)
records('Scenario Trials',d['trials'],p,'288 exploratory trial endpoints; source pair indices are preserved. No measured V995 performance.')
records('Scenario Contrasts',[dict(contrast=k,**r) for k,rr in d['contrasts'].items() for r in rr],p,'Conditional paired binary comparisons.')
records('Continuous Contrasts',[dict(contrast=k,**r) for k,rr in d['continuous_contrasts'].items() for r in rr],p,'Paired differences; min/max are not confidence bounds.')
metadata(p,{k:v for k,v in d.items() if k not in ['trials','contrasts','continuous_contrasts']})
p='Data/survivable_set_convergence.json';d=readj(p)
records('Convergence',d['trials'],p,'Historical dt diagnostic; does not establish model validity.')
metadata(p,{k:v for k,v in d.items() if k!='trials'})
p='Data/monte_carlo_results.json';d=readj(p)
records('Monte Carlo',[dict(scenario=s['label'],initial_rate_rad_s=float(k),**v) for s in d['scenarios'] for k,v in s['by_rate'].items()],p,'Committed summaries. Individual draws were not saved.')
for s in d['scenarios']:metadata(p,{'scenario.'+s['label']:{k:v for k,v in s.items() if k!='by_rate'}})
metadata(p,{k:v for k,v in d.items() if k!='scenarios'})
p='Data/release_recovery_results.json';d=readj(p)
records('Release Envelope',[dict(tier=k,**v) for k,v in d['envelope_by_tier'].items()],p,'Best-tier bound is sensor-limited, not a demonstrated dynamic failure boundary.')
metadata(p,{k:v for k,v in d.items() if k!='envelope_by_tier'})
records('Simulation Metadata',meta,'Data/*.json','Source registration, assumptions, seeds, reported verdicts and runtime metadata.')

from Analysis import recovery,guard,sim_release_recovery as sim
from Analysis.budget import rollup,load_mass_budget
recovery_records=[]
for c in recovery.default_cases():recovery_records.append(dict(**dataclasses.asdict(c),**{'result.'+k:v for k,v in recovery.evaluate_case(c).items()}))
records('Recovery Envelope Model',recovery_records,'Analysis/recovery.py','Recomputed first-order envelope. Historical defaults differ from current mass budget and 6-DoF model; intentionally preserved.')
records('Simulation Parameters',[dict(tier=n,**dataclasses.asdict(f())) for n,f in [('best',sim.best_params),('nominal',sim.nominal_params),('worst',sim.worst_params)]],'Analysis/sim_release_recovery.py','Historical modeled parameters; not V995 measurements.')
r=sim.simulate(sim.nominal_params(),2.,np.radians(60.))
expected=readj('Data/release_recovery_results.json')['demo_case']
assert r['recovery_time_s']==expected['recovery_time_s'] and r['max_descent_m']==expected['altitude_lost_m'] and r['peak_thrust_n']==expected['peak_thrust_n']
log=r['log']
add('Release Trace',[['time_s','tilt_rad','body_rate_rad_s','altitude_z_m','vertical_velocity_m_s','collective_thrust_N']]+[[float(log[k][i]) for k in ['t','tilt','omega_mag','z','vz','thrust']] for i in range(len(log['t']))],'Analysis/sim_release_recovery.py','REPRODUCED 2026-09-28. Full 12000 logged samples, dt=0.0005 s. World z positive upward; ground at -3 m. No downsampling.')

# Live, inspectable calculations linked to source tables.
mass=next(t for t in tabs if t['name']=='Mass Budget'); nr=len(mass['rows'])
add('Sizing Calculations',[
['Quantity','Value','Unit','Basis'],
['Best mass',f"=SUM('Mass Budget'!C2:C{nr})",'g','Historical design budget'],
['Nominal mass',f"=SUM('Mass Budget'!D2:D{nr})",'g','Historical design budget'],
['Worst mass',f"=SUM('Mass Budget'!E2:E{nr})",'g','Historical design budget'],
['Unmodeled hardware',5,'g','Selected allowance; not itemized'],
['Proposed frozen maximum','=CEILING(B4+B5,5)','g','Analysis/budget.py'],
['Nominal margin','=B6-B3','g','Frozen maximum minus nominal'],
['Abort threshold',225,'g','Selected design threshold; not a V995 limit'],
['Required thrust-to-weight',2,'ratio','REQ-PROP-001'],
['Number of motors',4,'count','Quadrotor'],
['Required per-motor thrust','=B9*B6/B10','gf','At proposed frozen maximum'],
['Required total thrust','=B9*B6/1000*9.81','N','g=9.81 m/s²'],
['5 V nominal sensor current',"=SUM('Power Budget'!C2:C5)",'mA','Listed 5 V loads only'],
['5 V worst sensor current',"=SUM('Power Budget'!D2:D5)",'mA','Listed 5 V loads only'],
['5 V nominal sensor power','=5*B13/1000','W','Excludes motors and unlisted loads'],
['Battery capacity',0.55,'Ah','Historical 2S pack example'],
['Usable capacity fraction',0.8,'fraction','Sizing assumption'],
['Average flight current',9,'A','Illustrative assumption, not measured'],
['Estimated flight duration','=60*B16*B17/B18','min','Capacity/current approximation'],
['Propeller diameter',50.8,'mm','Historical 2 inch example'],
['Radial clearance',2.5,'mm','Preliminary target'],
['Wall thickness',2,'mm','Preliminary target'],
['Guard inner diameter','=B20+2*B21','mm','Geometry example'],
['Guard outer diameter','=B23+2*B22','mm','Geometry example']
],'Analysis/budget.py; Design Report/calculations.md','Live formulas. Historical design only.',formulas=True)
rows=[['Case','Impact energy J','Stiffness N/m','Stress per force Pa/N','Clearance m','Allowable stress Pa','Deflection m','Equivalent force N','Impact stress Pa','Clearance SF','Stress SF','Functional pass']]
for i,c in enumerate(guard.default_cases(),2):
 rows.append([c.name,c.impact_energy_j,c.stiffness_n_m,c.stress_per_force_pa_n,c.clearance_m,c.elastic_allowable_pa,f'=SQRT(2*B{i}/C{i})',f'=SQRT(2*B{i}*C{i})',f'=D{i}*H{i}',f'=E{i}/G{i}',f'=F{i}/I{i}',f'=AND(J{i}>=2,K{i}>=3)'])
add('Guard Calculations',rows,'Analysis/guard.py','Assumed linear-elastic inputs. Does not establish fracture survival.',formulas=True)
rows=[['Collective fraction','Max total thrust N','Arm m','Collective thrust N','Differential per motor N','Roll-pitch authority Nm']]
for i in range(101):
 j=i+2;rows.append([i/100,4.2,0.06,f'=A{j}*B{j}',f'=MAX(0,MIN(D{j}/4,B{j}/4-D{j}/4))',f'=2*SQRT(2)*C{j}*E{j}'])
add('Mixer Authority',rows,'Analysis/sim_release_recovery.py','Four healthy symmetric rotors; assumed 60 mm arm. Not the rotor-out allocation solution.',formulas=True)
add('Statistics Calculations',[
 ['Calculation','n','Events or successes','Confidence','Result','Interpretation'],
 ['Zero-event upper bound',300,0,0.95,'=1-(1-D2)^(1/B2)','Exact one-sided upper event-rate bound'],
 ['Zero-event upper bound',1000,0,0.95,'=1-(1-D3)^(1/B3)','Exact one-sided upper event-rate bound'],
 ['Required n at max rate 0.01',None,0.01,0.95,'=CEILING(LN(1-D4)/LN(1-C4),1)','Zero observed events required'],
 ['Required n at max rate 0.003',None,0.003,0.95,'=CEILING(LN(1-D5)/LN(1-C5),1)','Zero observed events required'],
 ['Success fraction',300,0,0.95,'=C6/B6','Example matching zero-success primary cells'],
 ['One-sided success lower',300,0,0.95,'=IF(C7=0,0,BETA.INV(1-D7,C7,B7-C7+1))','Clopper-Pearson; sampling uncertainty only'],
 ['One-sided success upper',300,0,0.95,'=IF(C8=B8,1,BETA.INV(D8,C8+1,B8-C8))','Clopper-Pearson; each endpoint is one-sided 95%'],
 ],'Analysis/classifier_stats.py; Analysis/monte_carlo_recovery.py','Illustrative live formulas. No Tango interval or non-inferiority claim added.',formulas=True)

# Preserve original workbook cells and formulas. Prefix tabs and translate only sheet references.
wbpath='Research/Component Study/Drone_Component_and_Propulsion_Study.xlsx'
w=openpyxl.load_workbook(ROOT/wbpath)
mapping={s.title:'CS '+s.title.replace('Assumptions & Sources','Sources') for s in w}
for s in w:
 rows=[]
 for row in s:
  rr=[]
  for c in row:
   v=c.value
   if c.data_type=='f':
    for old,new in mapping.items():v=v.replace("'"+old+"'!","'"+new+"'!")
   rr.append(v)
  rows.append(rr)
 add(mapping[s.title],rows,wbpath+'#'+s.title,'Original component workbook; all 157 formulas retained with translated sheet references. Historical catalog assumptions.',header=0,formulas=True)
p='Research/Component Study/component-selection-2026/Component_Selection_2026.xlsx'
w=openpyxl.load_workbook(ROOT/p,data_only=False)
rows=[['Original sheet','Source row']+['Source column '+chr(65+i) for i in range(10)]]
for s in w:
 for i,row in enumerate(s.iter_rows(values_only=True),1):
  if any(v is not None for v in row):rows.append([s.title,i]+list(row))
add('Component Selection',rows,p,'Lossless nonempty-row archive of all 22 original tabs. Original headings retained as rows; source column letters map back to workbook.')
rr=[]
for p in sorted((ROOT/'Research/Component Study/unrestricted-tier-study-2026/results').glob('*.json')):
 for k,v in flatten(json.loads(p.read_text())).items():rr.append([p.stem,k,v])
add('Tier Research',[['Research section','Field','Value']]+rr,'Research/Component Study/unrestricted-tier-study-2026/results/','Catalog/research snapshot; nested candidate lists retained as JSON.')

# Math and software source archive: every top-level declaration is retained with line references.
decls=[];constants=[]
for p in sorted((ROOT/'Analysis').glob('*.py')):
 if p.name=='__init__.py':continue
 txt=p.read_text();lines=txt.splitlines();tree=ast.parse(txt)
 for node in tree.body:
  if isinstance(node,(ast.FunctionDef,ast.ClassDef)):
   code='\n'.join(lines[node.lineno-1:node.end_lineno])
   for part in range(0,len(code),35000):decls.append([p.name,node.name,node.lineno,part//35000+1,ast.get_docstring(node) or '',code[part:part+35000],link(p.relative_to(ROOT))+f'#L{node.lineno}'])
  elif isinstance(node,(ast.Assign,ast.AnnAssign)):
   constants.append([p.name,node.lineno,'\n'.join(lines[node.lineno-1:node.end_lineno]),link(p.relative_to(ROOT))+f'#L{node.lineno}'])
add('Model Functions',[['File','Function or class','Source line','Part','Description','Python definition','Source URL']]+decls,'Analysis/*.py','Exact pinned implementation, including controllers, allocator, event handling, statistics and plotting. Source text is not executed by Sheets.')
add('Model Constants',[['File','Source line','Declaration','Source URL']]+constants,'Analysis/*.py','Complete module-level declarations; expressions preserved as source text.')
eq=[
 ['Mass freeze','m_frozen = 5 ceil((sum(m_worst)+5)/5)','g','Sizing Calculations','Analysis/budget.py','5 g allowance is selected, not measured'],
 ['Thrust requirement','T_motor = (T/W) m / 4','gf with mass in g','Sizing Calculations','Analysis/gates.py','Historical quadrotor sizing'],
 ['Flight time','t_min = 60 C_Ah f_usable / I_A','min','Sizing Calculations','Design Report/calculations.md','Constant average current approximation'],
 ['Guard geometry','D_inner=D_prop+2c; D_outer=D_inner+2t','mm','Sizing Calculations','Design Report/calculations.md','Clearance and wall are targets'],
 ['Angular arrest','t = I |omega| / tau','s','Recovery Envelope Model','Analysis/recovery.py','Constant opposing torque'],
 ['Rest-to-rest rotation','t = 2 sqrt(|theta| I / tau)','s','Recovery Envelope Model','Analysis/recovery.py','Bang-bang single-axis model'],
 ['Vertical envelope','a_down = g - (T_up + 0.5 rho Cd A v_down^2)/m','m/s²','Recovery Envelope Model','Analysis/recovery.py','Downward-positive convention'],
 ['Rigid-body dynamics','omega_dot = I^-1 (tau - omega cross (I omega))','rad/s²','Model Functions','Analysis/rigid_body.py','Body angular rates, diagonal inertia'],
 ['Attitude kinematics','q_dot = 0.5 q tensor (0,omega); normalize(q_next)','1/s','Model Functions','Analysis/rigid_body.py','Quaternion w,x,y,z; body to world'],
 ['6-DoF translation','m v_dot = R(q)[0,0,T] + F_drag - [0,0,mg]','N','Model Functions','Analysis/sim_release_recovery.py','World z positive upward; see exact integrator'],
 ['Mixer differential','d=max(0,min(T/4,T_max/4-T/4))','N','Mixer Authority','Analysis/sim_release_recovery.py','Four healthy motors'],
 ['Mixer roll/pitch bound','tau=2 sqrt(2) arm d','N m','Mixer Authority','Analysis/sim_release_recovery.py','Bound at current collective'],
 ['Healthy trim','sum(f)=mg; B_moment f + tau_CG + bias = 0; 0<=f<=cap','N, N m','Model Functions','Analysis/failure_allocation.py','Infeasible trim remains a reported outcome'],
 ['Fault allocation','wrench=B f; bounded weighted allocation, torque pass then collective','N, N m','Model Functions','Analysis/failure_allocation.py','See exact apply() implementation; no replacement optimizer'],
 ['Yaw drag','tau_z=-c omega_z |omega_z|','N m','Model Functions','Analysis/sim_release_recovery.py','A2 coefficient assumed'],
 ['Guard deflection','delta=sqrt(2E/k)','m','Guard Calculations','Analysis/guard.py','Linear elastic energy model'],
 ['Guard force','F=sqrt(2Ek)','N','Guard Calculations','Analysis/guard.py','Equivalent elastic force, not a measured peak'],
 ['Guard stress','sigma=c_sigma F','Pa','Guard Calculations','Analysis/guard.py','Material/geometry assumptions'],
 ['Guard pass','clearance/delta >= 2 AND allowable/sigma >= 3','dimensionless','Guard Calculations','Analysis/guard.py','Functional check only'],
 ['Zero-event bound','p_upper=1-(1-confidence)^(1/n)','fraction','Statistics Calculations','Analysis/classifier_stats.py','Independent Bernoulli trials'],
 ['Zero-event sample size','n=ceil(ln(1-confidence)/ln(1-p_max))','count','Statistics Calculations','Analysis/classifier_stats.py','Design calculation'],
 ['Binomial bounds','lower=BetaInv(alpha,s,n-s+1); upper=BetaInv(1-alpha,s+1,n-s)','fraction','Statistics Calculations','Analysis/monte_carlo_recovery.py','Boundary cases 0 and n handled explicitly'],
 ['Paired difference','delta=(only_a-only_b)/n','fraction','Paired A; Paired A2','Analysis/survivable_set.py','Paired discordance, not independent groups'],
 ['Landing proxy','abs(vz_contact)<=v_limit AND tilt_contact<=tilt_limit','m/s, rad','Study A; Study A2','Analysis/survivable_set.py','Assumed thresholds; no physical validation'],
 ['Parachute descent','Ballistic fall through deployment+inflation, then quadratic drag toward terminal speed','m/s','Model Functions','Analysis/survivable_set.py','Model-to-model calibration; no V995 drop data'],
 ['Free fall','d=0.5 g t^2; v=g t','m, m/s','Calculation Notes','Design Report/calculations.md','From rest, no drag'],
]
add('Equation Register',[['Calculation','Equation or method','Units','Workbook location','Implementation','Limitations']]+eq,'Analysis/*.py; Design Report/calculations.md','Readable equation map; full exact implementation retained in Model Functions and Model Constants.')

notes=[]
paths=['Design Report/calculations.md','Analysis/current-results.md','docs/specs/survivable-set/design.md','docs/specs/survivable-set/design-a2.md','docs/specs/survivable-set/scenario-contract.md','docs/bench-acquisition.md','docs/data-and-figures.md','docs/specs/measured-authority-gate/evidence-contract.md','Engineering Data/component_specs_under_load.md','literature/claim-ledger.md']
for p in paths:
 text=(ROOT/p).read_text()
 section='Introduction'
 for i,paragraph in enumerate(text.split('\n\n'),1):
  if paragraph.startswith('#'):section=paragraph.splitlines()[0].lstrip('# ')
  if paragraph.strip():notes.append([p,section,i,paragraph,link(p)])
add('Calculation Notes',[['Source','Section','Block','Original text','Source URL']]+notes,'Design and evidence documentation','Complete selected calculation/methodology documents. Original prose retained, including uncertainty and unresolved decisions.')
figs=[]
for p in sorted((ROOT/'Figures').iterdir()):
 figs.append([p.name,p.suffix[1:],link(p.relative_to(ROOT)),'Release Trace / Release Envelope' if 'release_recovery' in p.name else 'Study A / Study A2 / Package Rankings' if 'survivable_set' in p.name else 'Explanatory diagram'])
add('Figure Index',[['Figure','Format','Pinned source URL','Underlying data']]+figs,'Figures/','Original figures linked alongside native workbook charts. No image values were digitized.')

# Chart-ready summaries, retaining cohort labels and avoiding pooled inference.
trialtab=next(t for t in tabs if t['name']=='Scenario Trials')
head=trialtab['rows'][0];trials=[dict(zip(head,row)) for row in trialtab['rows'][1:]]
summary=[]
for cl,cell,package,arm in sorted(set((r['class'],r['cell'],r['package'],r['arm']) for r in trials)):
 rr=[r for r in trials if (r['class'],r['cell'],r['package'],r['arm'])==(cl,cell,package,arm)]
 summary.append(dict(failure_class=cl,cell=cell,package=package,arm=arm,n=len(rr),safe_count=sum(r['safe'] for r in rr),mean_impact_speed_m_s=sum(r['impact_speed_m_s'] for r in rr)/len(rr),mean_contact_time_s=sum(r['t_contact_s'] for r in rr)/len(rr)))
records('Scenario Summary',summary,'Data/scenario-diagnostic/scenario_diagnostic.json','Derived descriptive means by exact cohort. No pooled success inference.')

for t in tabs:
 if t is sources:continue
 sources['rows'].append([t['name'],t['source'],link(t['source']) if t['source'] and not any(c in t['source'] for c in ['*',';','#']) else '',len(t['rows'])-t['header'],t['note']])
sources['width']=5
out=BASE/'payload';out.mkdir(exist_ok=True)
manifest=[]
for t in tabs:
 rows=t.pop('rows');t['nrows']=len(rows);t['rowCount']=max(100,t['nrows']+10,250 if t['name']=='Charts' else 0);t['columnCount']=max(t['width'],12)
 t['chunks']=[]
 for start in range(0,len(rows),600):
  fn=f"{t['id']}_{start}.json"; (out/fn).write_text(json.dumps(rows[start:start+600],ensure_ascii=False,allow_nan=False));t['chunks'].append({'file':fn,'start':start,'count':len(rows[start:start+600])})
 manifest.append(t)
(out/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False))
print(json.dumps({'tabs':len(tabs),'rows':sum(t['nrows'] for t in tabs),'formula_tabs':[t['name'] for t in tabs if t['formulas']],'sizes':[(t['name'],t['nrows'],t['width']) for t in tabs]},indent=2))
